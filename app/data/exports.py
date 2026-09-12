"""Downloads: CSVs and a ready-to-work Excel workbook.

Prices only - never returns. Computing the return series at three frequencies
is Module 2's exercise, and handing it over would remove the point.
"""
from __future__ import annotations

import io

import pandas as pd

from .datasets import Dataset


def _stamp(ds: Dataset) -> str:
    return f"{ds.name}_{ds.summary()['last_date']}"


def prices_csv(ds: Dataset) -> bytes:
    df = ds.prices.copy()
    df.index.name = "Date"
    return df.round(4).to_csv().encode()


def benchmarks_csv(ds: Dataset) -> bytes:
    df = ds.benchmarks.copy()
    df.index.name = "Date"
    return df.round(4).to_csv().encode()


def universe_csv(ds: Dataset) -> bytes:
    cols = ["epic", "yahoo", "name", "sector", "industry",
            "mcap_gbp_m", "weight_pct", "first_date", "n_obs", "currency"]
    df = ds.universe[[c for c in cols if c in ds.universe.columns]].copy()
    return df.to_csv(index=False).encode()


def quality_csv(ds: Dataset) -> bytes:
    """The cleaning log - published so joiners can check their own outlier hunt."""
    rows = []
    for rep in ds.quality_report:
        for kind in ("spikes", "breaks"):
            for rec in rep[kind]:
                rows.append({"ticker": rep["ticker"],
                             "issue": "repaired_bad_print" if kind == "spikes" else "scale_break_excluded",
                             **{k: v for k, v in rec.items() if k != "bad_dates"}})
    for t in ds.excluded.get("data_quality", []):
        rows.append({"ticker": t, "issue": "EXCLUDED_from_dataset", "date": "", "return_pct": ""})
    for t in ds.excluded.get("short_history", []):
        rows.append({"ticker": t, "issue": "EXCLUDED_insufficient_history", "date": "", "return_pct": ""})
    if not rows:
        return b"ticker,issue\n"
    return pd.DataFrame(rows).to_csv(index=False).encode()


def workbook(ds: Dataset) -> bytes:
    """Starter .xlsx: the same data on three tabs, plus a README tab.

    Saves the CSV-import faff so the joiner gets to the maths sooner. It does
    NOT contain any returns, covariance or formulas - that is the exercise.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.dataframe import dataframe_to_rows

    s = ds.summary()
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="EFE7DC")
    head_font = Font(bold=True, color="3A3340")

    # --- README -----------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    lines = [
        ("FTSE 100 Risk Model Training - starter data", True),
        ("", False),
        (f"Dataset      {s['label']}", False),
        (f"Window       {s['first_date']} to {s['last_date']}  ({s['n_days']} trading days)", False),
        (f"Names        {s['n_stocks']}", False),
        (f"Generated    {s['last_date']}", False),
        ("", False),
        ("Tabs", True),
        ("  Prices       Adjusted closing prices in PENCE (GBp). Dividends and", False),
        ("               splits are already reflected, so price changes ARE", False),
        ("               total returns for each stock.", False),
        ("  Benchmarks   Index levels. Read the notes - they do not all mean", False),
        ("               the same thing, and two of them disagree on purpose.", False),
        ("  Universe     Name, ICB industry, market cap and index weight.", False),
        ("", False),
        ("What is deliberately NOT here", True),
        ("  Returns, covariances and volatilities. Building those is the", False),
        ("  exercise. Start with Module 2 on the website.", False),
        ("", False),
        ("Health warnings", True),
        ("  Survivorship bias: these are TODAY'S index members. Names that", False),
        ("  dropped out of the FTSE 100, often after doing badly, are absent.", False),
        ("  Your volatility estimates will be flattered.", False),
        ("", False),
        ("  Cleaning: bad prints have been repaired and two names excluded.", False),
        ("  The full log is on the website - but try finding them yourself first.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if bold:
            c.font = Font(bold=True, size=12 if i == 1 else 11, color="3A3340")
    ws.column_dimensions["A"].width = 76

    # --- data tabs --------------------------------------------------------
    def add(df: pd.DataFrame, title: str, index_name: str | None):
        sheet = wb.create_sheet(title)
        out = df.copy()
        if index_name:
            out.index.name = index_name
            out = out.reset_index()
            if index_name == "Date":
                out["Date"] = pd.to_datetime(out["Date"]).dt.date
        for row in dataframe_to_rows(out, index=False, header=True):
            sheet.append(row)
        for cell in sheet[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 12
        for j in range(2, min(out.shape[1], 200) + 1):
            sheet.column_dimensions[get_column_letter(j)].width = 11
        return sheet

    add(ds.prices.round(4), "Prices", "Date")
    add(ds.benchmarks.round(4), "Benchmarks", "Date")
    cols = ["epic", "yahoo", "name", "sector", "industry", "mcap_gbp_m", "weight_pct", "first_date"]
    add(ds.universe[[c for c in cols if c in ds.universe.columns]], "Universe", None)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


FILES = {
    "prices": ("prices", "csv", prices_csv, "Adjusted daily prices in pence, one column per stock"),
    "benchmarks": ("benchmarks", "csv", benchmarks_csv, "Index levels: price, total return, and two composites"),
    "universe": ("universe", "csv", universe_csv, "Name, ICB industry, market cap and index weight"),
    "quality": ("cleaning-log", "csv", quality_csv, "Every repair and exclusion we made"),
    "workbook": ("starter", "xlsx", workbook, "All three tabs in one Excel workbook, ready to work in"),
}


def filename(key: str, ds: Dataset) -> str:
    base, ext, *_ = FILES[key]
    return f"ftse100-{base}-{_stamp(ds)}.{ext}"
